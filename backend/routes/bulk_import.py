import csv
import io
import json
import re
import uuid
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook

from models import ROIRecord
from services.storage import delete_by_batch_id, get_all_records, save_record

router = APIRouter(prefix="/api")


class SkipRow(Exception):
    """Row has partial data but can't be imported — caller should flag for review."""
    pass

# Maps spreadsheet column headers (lowercased) to ROIRecord field names.
# Covers both the user-facing labels from the sample file and the raw backend keys.
_COLUMN_MAP = {
    "year":                           "year",
    "client":                         "client",
    "publisher":                      "publisher",
    "currency":                       "currency",
    "date delivered":                 "date_delivered",
    "date_delivered":                 "date_delivered",
    "identified risk":                "identified_risk",
    "identified_risk":                "identified_risk",
    "identified cost avoidance":      "id_cost_avoidance",
    "id_cost_avoidance":              "id_cost_avoidance",
    "accomplished cost avoidance":    "acc_cost_avoidance",
    "acc_cost_avoidance":             "acc_cost_avoidance",
    "identified cost optimization":   "id_cost_optimization",
    "id_cost_optimization":           "id_cost_optimization",
    "accomplished cost optimization": "acc_cost_optimization",
    "acc_cost_optimization":          "acc_cost_optimization",
    "realized cost savings":          "realized_savings",
    "realized_savings":               "realized_savings",
    "annual publisher contract":      "contract_spend",
    "contract_spend":                 "contract_spend",
    "month":                          "month",
    "notes":                          "notes",
    "sme":                            "sme",
    "pricing available":              "pricing_available",
    "pricing_available":              "pricing_available",
}

_NUMERIC_FIELDS = {
    "identified_risk", "id_cost_avoidance", "acc_cost_avoidance",
    "id_cost_optimization", "acc_cost_optimization", "realized_savings",
    "contract_spend",
}

# Canonical publisher names keyed by their lowercased variants.
# Covers company names, common mis-spellings, abbreviations, and product names
# that are often used in place of the publisher (e.g. "Java" → Oracle).
_PUBLISHER_NORMS: dict[str, str] = {
    # ── Microsoft ────────────────────────────────────────────────────────────
    "microsoft":                "Microsoft",
    "ms":                       "Microsoft",
    "msft":                     "Microsoft",
    "windows":                  "Microsoft",
    "office":                   "Microsoft",
    "office 365":               "Microsoft",
    "o365":                     "Microsoft",
    "microsoft 365":            "Microsoft",
    "m365":                     "Microsoft",
    "sql server":               "Microsoft",
    "microsoft sql server":     "Microsoft",
    "mssql":                    "Microsoft",
    "exchange":                 "Microsoft",
    "sharepoint":               "Microsoft",
    "teams":                    "Microsoft",
    "microsoft teams":          "Microsoft",
    "dynamics":                 "Microsoft",
    "dynamics 365":             "Microsoft",
    "power bi":                 "Microsoft",
    "power platform":           "Microsoft",
    "power apps":               "Microsoft",
    "power automate":           "Microsoft",
    "github":                   "Microsoft",
    "visual studio":            "Microsoft",
    "intune":                   "Microsoft",
    "sccm":                     "Microsoft",
    "system center":            "Microsoft",
    "active directory":         "Microsoft",
    "hyper-v":                  "Microsoft",
    "skype":                    "Microsoft",
    "onedrive":                 "Microsoft",
    # ── Azure (tracked separately from Microsoft in many engagements) ────────
    "azure":                    "Azure",
    "microsoft azure":          "Azure",
    "azure ad":                 "Azure",
    "azure active directory":   "Azure",
    # ── Oracle ───────────────────────────────────────────────────────────────
    "oracle":                   "Oracle",
    "java":                     "Oracle",
    "jdk":                      "Oracle",
    "jre":                      "Oracle",
    "mysql":                    "Oracle",
    "peoplesoft":               "Oracle",
    "jd edwards":               "Oracle",
    "jde":                      "Oracle",
    "j.d. edwards":             "Oracle",
    "siebel":                   "Oracle",
    "hyperion":                 "Oracle",
    "netsuite":                 "Oracle",
    "oracle database":          "Oracle",
    "oracle db":                "Oracle",
    "oracle rdbms":             "Oracle",
    "weblogic":                 "Oracle",
    "goldengate":               "Oracle",
    "solaris":                  "Oracle",
    "sparc":                    "Oracle",
    "sun microsystems":         "Oracle",
    "primavera":                "Oracle",
    "agile plm":                "Oracle",
    "oracle agile":             "Oracle",
    "eloqua":                   "Oracle",
    "oracle eloqua":            "Oracle",
    "oci":                      "Oracle",
    "oracle cloud":             "Oracle",
    "oracle cloud infrastructure": "Oracle",
    "virtualbox":               "Oracle",
    # ── IBM ──────────────────────────────────────────────────────────────────
    "ibm":                      "IBM",
    "db2":                      "IBM",
    "ibm db2":                  "IBM",
    "websphere":                "IBM",
    "cognos":                   "IBM",
    "ibm cognos":               "IBM",
    "spss":                     "IBM",
    "ibm spss":                 "IBM",
    "rational":                 "IBM",
    "tivoli":                   "IBM",
    "maas360":                  "IBM",
    "ibm maas360":              "IBM",
    "qradar":                   "IBM",
    "ibm qradar":               "IBM",
    "turbonomic":               "IBM",
    "infosphere":                "IBM",
    "datastage":                "IBM",
    "ibm datastage":            "IBM",
    "sterling":                 "IBM",
    "ibm sterling":             "IBM",
    "planning analytics":       "IBM",
    "watson":                   "IBM",
    "ibm watson":               "IBM",
    "instana":                  "IBM",
    # ── Red Hat (IBM subsidiary, often tracked separately) ───────────────────
    "red hat":                  "Red Hat",
    "redhat":                   "Red Hat",
    "rhel":                     "Red Hat",
    "red hat enterprise linux": "Red Hat",
    "openshift":                "Red Hat",
    "red hat openshift":        "Red Hat",
    "ansible":                  "Red Hat",
    "red hat ansible":          "Red Hat",
    "jboss":                    "Red Hat",
    # ── SAP ──────────────────────────────────────────────────────────────────
    "sap":                      "SAP",
    "sap se":                   "SAP",
    "sap hana":                 "SAP",
    "hana":                     "SAP",
    "s/4hana":                  "SAP",
    "s4hana":                   "SAP",
    "successfactors":           "SAP",
    "sap successfactors":       "SAP",
    "ariba":                    "SAP",
    "sap ariba":                "SAP",
    "concur":                   "SAP",
    "sap concur":               "SAP",
    "business objects":         "SAP",
    "businessobjects":          "SAP",
    "crystal reports":          "SAP",
    "sap business one":         "SAP",
    "sap erp":                  "SAP",
    "sap bw":                   "SAP",
    "sap crm":                  "SAP",
    # ── VMware (Broadcom) ────────────────────────────────────────────────────
    "vmware":                   "VMware",
    "vm ware":                  "VMware",
    "vm-ware":                  "VMware",
    "broadcom vmware":          "VMware",
    "vmware by broadcom":       "VMware",
    "vsphere":                  "VMware",
    "vmware vsphere":           "VMware",
    "esxi":                     "VMware",
    "vmware esxi":              "VMware",
    "vcenter":                  "VMware",
    "nsx":                      "VMware",
    "vmware nsx":               "VMware",
    "horizon":                  "VMware",
    "vmware horizon":           "VMware",
    "carbon black":             "VMware",
    "vmware carbon black":      "VMware",
    "tanzu":                    "VMware",
    "vsan":                     "VMware",
    "vrealize":                 "VMware",
    "aria":                     "VMware",
    "workspace one":            "VMware",
    "workspace 1":              "VMware",
    # ── Broadcom (non-VMware) ────────────────────────────────────────────────
    "broadcom":                 "Broadcom",
    "symantec":                 "Broadcom",
    "ca technologies":          "Broadcom",
    "ca":                       "Broadcom",
    "brocade":                  "Broadcom",
    # ── Salesforce ───────────────────────────────────────────────────────────
    "salesforce":               "Salesforce",
    "mulesoft":                 "Salesforce",
    "salesforce mulesoft":      "Salesforce",
    "tableau":                  "Salesforce",
    "salesforce tableau":       "Salesforce",
    "slack":                    "Salesforce",
    "heroku":                   "Salesforce",
    "pardot":                   "Salesforce",
    "marketing cloud":          "Salesforce",
    "salesforce marketing cloud": "Salesforce",
    "commerce cloud":           "Salesforce",
    "einstein":                 "Salesforce",
    "quip":                     "Salesforce",
    # ── ServiceNow ───────────────────────────────────────────────────────────
    "servicenow":               "ServiceNow",
    "service now":              "ServiceNow",
    # ── Adobe ────────────────────────────────────────────────────────────────
    "adobe":                    "Adobe",
    "creative cloud":           "Adobe",
    "adobe creative cloud":     "Adobe",
    "acrobat":                  "Adobe",
    "adobe acrobat":            "Adobe",
    "marketo":                  "Adobe",
    "adobe marketo":            "Adobe",
    "experience manager":       "Adobe",
    "aem":                      "Adobe",
    "adobe analytics":          "Adobe",
    "adobe target":             "Adobe",
    "adobe campaign":           "Adobe",
    "photoshop":                "Adobe",
    "illustrator":              "Adobe",
    "indesign":                 "Adobe",
    "premiere pro":             "Adobe",
    "after effects":            "Adobe",
    # ── Cisco ────────────────────────────────────────────────────────────────
    "cisco":                    "Cisco",
    "cisco systems":            "Cisco",
    "webex":                    "Cisco",
    "cisco webex":              "Cisco",
    "meraki":                   "Cisco",
    "cisco meraki":             "Cisco",
    "appdynamics":              "Cisco",
    "cisco appdynamics":        "Cisco",
    "duo":                      "Cisco",
    "cisco duo":                "Cisco",
    "umbrella":                 "Cisco",
    "cisco umbrella":           "Cisco",
    "splunk":                   "Cisco",
    # ── AWS ──────────────────────────────────────────────────────────────────
    "aws":                      "AWS",
    "amazon web services":      "AWS",
    "amazon":                   "AWS",
    # ── Google Cloud ─────────────────────────────────────────────────────────
    "google":                   "Google Cloud",
    "google cloud":             "Google Cloud",
    "gcp":                      "Google Cloud",
    "google workspace":         "Google Cloud",
    "g suite":                  "Google Cloud",
    "bigquery":                 "Google Cloud",
    "looker":                   "Google Cloud",
    "apigee":                   "Google Cloud",
    # ── Workday ──────────────────────────────────────────────────────────────
    "workday":                  "Workday",
    "workday hcm":              "Workday",
    "workday financials":       "Workday",
    "adaptive":                 "Workday",
    "adaptive planning":        "Workday",
    "adaptive insights":        "Workday",
    # ── Autodesk ─────────────────────────────────────────────────────────────
    "autodesk":                 "Autodesk",
    "autocad":                  "Autodesk",
    "revit":                    "Autodesk",
    "inventor":                 "Autodesk",
    "3ds max":                  "Autodesk",
    "maya":                     "Autodesk",
    "fusion 360":               "Autodesk",
    "civil 3d":                 "Autodesk",
    "navisworks":               "Autodesk",
    "bim 360":                  "Autodesk",
    # ── Atlassian ────────────────────────────────────────────────────────────
    "atlassian":                "Atlassian",
    "jira":                     "Atlassian",
    "confluence":               "Atlassian",
    "bitbucket":                "Atlassian",
    "trello":                   "Atlassian",
    "opsgenie":                 "Atlassian",
    # ── OpenText (formerly Micro Focus) ──────────────────────────────────────
    "opentext":                 "OpenText",
    "open text":                "OpenText",
    "micro focus":              "OpenText",
    "microfocus":               "OpenText",
    "cobol":                    "OpenText",
    "alm":                      "OpenText",
    "loadrunner":               "OpenText",
    "documentum":               "OpenText",
    # ── BMC Software ─────────────────────────────────────────────────────────
    "bmc":                      "BMC",
    "bmc software":             "BMC",
    "control-m":                "BMC",
    "control m":                "BMC",
    "remedy":                   "BMC",
    "helix":                    "BMC",
    "truesight":                "BMC",
    # ── Pegasystems ──────────────────────────────────────────────────────────
    "pegasystems":              "Pegasystems",
    "pega":                     "Pegasystems",
    # ── Informatica ──────────────────────────────────────────────────────────
    "informatica":              "Informatica",
    "powercenter":              "Informatica",
    "iics":                     "Informatica",
    # ── PTC ──────────────────────────────────────────────────────────────────
    "ptc":                      "PTC",
    "windchill":                "PTC",
    "thingworx":                "PTC",
    "creo":                     "PTC",
    "vuforia":                  "PTC",
    # ── Citrix ───────────────────────────────────────────────────────────────
    "citrix":                   "Citrix",
    "xenapp":                   "Citrix",
    "xendesktop":               "Citrix",
    "citrix virtual apps":      "Citrix",
    # ── Nutanix ──────────────────────────────────────────────────────────────
    "nutanix":                  "Nutanix",
    # ── Veritas ──────────────────────────────────────────────────────────────
    "veritas":                  "Veritas",
    # ── SolarWinds ───────────────────────────────────────────────────────────
    "solarwinds":               "SolarWinds",
    "solar winds":              "SolarWinds",
}


def _normalize_publisher(name: str) -> str:
    """Return the canonical form of a publisher name, or the original if unknown."""
    return _PUBLISHER_NORMS.get(name.strip().lower(), name.strip())


def _parse_number(val) -> float | None:
    """Return a float from a cell value, or None for blanks and non-numeric text."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    text = str(val).strip()
    if not text or text in ("-", "—"):
        return None
    cleaned = re.sub(r"[$,\s]", "", text)
    try:
        return float(cleaned)
    except ValueError:
        return None


def _parse_year(val) -> int | None:
    if val is None:
        return None
    if isinstance(val, int):
        return val
    try:
        return int(str(val).strip())
    except ValueError:
        return None


def _rows_from_xlsx(content: bytes) -> list[tuple[str, list[str], list[list]]]:
    """Return one entry per sheet: (sheet_name, headers, data_rows).
    The sheet name is the authoritative client name for all rows on that tab."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
        data_rows = [list(r) for r in all_rows[1:]]
        sheets.append((ws.title, headers, data_rows))
    wb.close()
    return sheets


def _rows_from_csv(content: bytes) -> tuple[list[str], list[list]]:
    text = content.decode("utf-8-sig")  # strip BOM if present
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
    return [h.strip() for h in rows[0]], rows[1:]


def _map_row(
    headers: list[str],
    values: list,
    client_override: str | None = None,
    batch_id: str | None = None,
) -> ROIRecord | None:
    """Map one spreadsheet row to an ROIRecord.

    Returns None for genuinely blank rows (no flagging needed).
    Raises SkipRow for rows with partial data that can't be saved — caller
    should collect these and surface them to the user for review.

    client_override: when provided (e.g. the Excel tab name), it replaces whatever
    is in the 'Client' column.
    batch_id: tags the record so the whole import can be undone in one call."""
    col_index: dict[int, str] = {}
    for i, h in enumerate(headers):
        key = _COLUMN_MAP.get(h.strip().lower())
        if key:
            col_index[i] = key

    raw: dict[str, object] = {}
    for i, key in col_index.items():
        raw[key] = values[i] if i < len(values) else None

    # Truly blank row — silently skip, no need to flag.
    if all(v is None or str(v).strip() == "" for v in raw.values()):
        return None

    year_val = _parse_year(raw.get("year"))
    if not year_val:
        raw_year = raw.get("year")
        raise SkipRow(
            f"Could not parse year{f': {raw_year!r}' if raw_year is not None else ' (column missing or empty)'}"
        )

    # Tab name is authoritative; fall back to the column value only for CSV.
    client = client_override if client_override else str(raw.get("client") or "").strip()
    publisher = _normalize_publisher(str(raw.get("publisher") or ""))
    if not publisher:
        raise SkipRow("Missing publisher")

    fields: dict = {
        "year":           year_val,
        "client":         client,
        "publisher":      publisher,
        "currency":       str(raw.get("currency") or "USD").strip(),
        "date_delivered": str(raw.get("date_delivered") or "").strip() or None,
        "month":          str(raw.get("month") or "").strip() or None,
        "notes":          str(raw.get("notes") or "").strip() or None,
        "sme":            str(raw.get("sme") or "").strip() or None,
        "source_file":    f"{client} — {publisher} — {year_val} (bulk import)",
        "batch_id":       batch_id,
    }
    for num_key in _NUMERIC_FIELDS:
        fields[num_key] = _parse_number(raw.get(num_key))

    return ROIRecord(**fields)


@router.post("/bulk-import")
async def bulk_import(file: UploadFile = File(...)):
    """Accept an XLSX or CSV file and stream SSE progress events while processing.

    Each event is a JSON object on a `data:` line. The final event has type "done"
    and carries the full result. The client reads the response body progressively."""
    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx", ".csv"):
        raise HTTPException(status_code=415, detail="Only .xlsx and .csv files are accepted.")

    content = await file.read()
    batch_id = f"batch_{uuid.uuid4().hex[:16]}"

    def generate():
        imported = 0
        errors: list[str] = []
        flagged: list[dict] = []
        sheets: list[dict] = []
        seen_keys: set[tuple] = set()

        # Snapshot of existing records keyed by (client, publisher, year) for
        # duplicate detection. Built once so row processing stays fast.
        existing_db: dict[tuple, dict] = {}
        for r in get_all_records():
            c = str(r.get("client") or "").strip().lower()
            p = str(r.get("publisher") or "").strip().lower()
            y = r.get("year")
            if c and p and y:
                existing_db[(c, p, int(y))] = r

        def sse(data: dict) -> str:
            return f"data: {json.dumps(data)}\n\n"

        def _row_data(headers: list[str], row: list, client_override: str | None) -> dict:
            data = {}
            for i, h in enumerate(headers):
                key = _COLUMN_MAP.get(h.strip().lower())
                if key:
                    val = row[i] if i < len(row) else None
                    data[key] = str(val).strip() if val is not None else ""
            if client_override:
                data["client"] = client_override
            return data

        def _process_row(row_num: int, row: list, headers: list[str],
                         client_override: str | None, location: str) -> str:
            nonlocal imported
            raw_data = _row_data(headers, row, client_override)
            try:
                record = _map_row(headers, row, client_override=client_override, batch_id=batch_id)
                if record is None:
                    return "blank"
                dup_key = (record.client.lower(), record.publisher.lower(), record.year)
                # Within-file duplicate
                if dup_key in seen_keys:
                    flagged.append({
                        "location": location, "row": row_num,
                        "reason": f"Duplicate within file — {record.client} / {record.publisher} / {record.year} already appears in this import",
                        "data": raw_data,
                    })
                    return "flagged"
                # Duplicate against existing database records
                if dup_key in existing_db:
                    existing = existing_db[dup_key]
                    flagged.append({
                        "location": location, "row": row_num,
                        "reason": f"Already in database — {record.client} / {record.publisher} / {record.year} was previously imported",
                        "data": raw_data,
                        "existing": {
                            "record_id":    existing.get("record_id"),
                            "saved_at":     existing.get("saved_at"),
                            "source_file":  existing.get("source_file"),
                            "sme":          existing.get("sme"),
                            "identified_risk":   existing.get("identified_risk"),
                            "acc_cost_avoidance": existing.get("acc_cost_avoidance"),
                            "realized_savings":  existing.get("realized_savings"),
                        },
                    })
                    return "flagged"
                seen_keys.add(dup_key)
                save_record(record)
                imported += 1
                return "imported"
            except SkipRow as e:
                flagged.append({"location": location, "row": row_num, "reason": str(e), "data": raw_data})
                return "flagged"
            except Exception as e:
                errors.append(f"{location} row {row_num}: {e}")
                return "error"

        try:
            if ext == ".xlsx":
                try:
                    sheet_data = _rows_from_xlsx(content)
                except Exception as e:
                    yield sse({"type": "error", "detail": f"Could not parse file: {e}"})
                    return
                if not sheet_data:
                    yield sse({"type": "error", "detail": "File appears to be empty."})
                    return

                yield sse({"type": "start", "format": "xlsx", "sheet_count": len(sheet_data)})

                for si, (sheet_name, headers, data_rows) in enumerate(sheet_data):
                    n = len(data_rows)
                    s_imported = s_flagged = 0
                    # Emit ~15 progress ticks per sheet regardless of size
                    interval = max(1, n // 15)

                    yield sse({
                        "type": "sheet_start",
                        "name": sheet_name,
                        "sheet_index": si + 1,
                        "total_sheets": len(sheet_data),
                        "row_count": n,
                    })

                    for row_num, row in enumerate(data_rows, start=2):
                        outcome = _process_row(row_num, row, headers, sheet_name, sheet_name)
                        if outcome == "imported":
                            s_imported += 1
                        elif outcome == "flagged":
                            s_flagged += 1

                        processed = row_num - 1
                        if processed % interval == 0 or processed == n:
                            yield sse({
                                "type": "progress",
                                "sheet": sheet_name,
                                "sheet_index": si + 1,
                                "total_sheets": len(sheet_data),
                                "processed": processed,
                                "total": n,
                            })

                    sheets.append({"name": sheet_name, "imported": s_imported, "flagged": s_flagged})
                    yield sse({"type": "sheet_done", "name": sheet_name, "imported": s_imported, "flagged": s_flagged})

            else:  # CSV
                try:
                    headers, data_rows = _rows_from_csv(content)
                except Exception as e:
                    yield sse({"type": "error", "detail": f"Could not parse file: {e}"})
                    return
                if not headers:
                    yield sse({"type": "error", "detail": "File appears to be empty."})
                    return

                n = len(data_rows)
                interval = max(1, n // 15)
                yield sse({"type": "start", "format": "csv", "row_count": n})

                for row_num, row in enumerate(data_rows, start=2):
                    _process_row(row_num, row, headers, None, "CSV")
                    processed = row_num - 1
                    if processed % interval == 0 or processed == n:
                        yield sse({"type": "progress", "processed": processed, "total": n})

        except Exception as e:
            yield sse({"type": "error", "detail": f"Processing failed: {e}"})
            return

        result: dict = {"type": "done", "batch_id": batch_id, "imported": imported, "flagged": flagged, "errors": errors}
        if sheets:
            result["sheets"] = sheets
        yield sse(result)

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.delete("/bulk-import/{batch_id}")
def undo_bulk_import(batch_id: str):
    """Delete all records from a specific bulk import batch (undo a bulk import)."""
    deleted = delete_by_batch_id(batch_id)
    if deleted == 0:
        raise HTTPException(status_code=404, detail="No records found for this batch.")
    return {"deleted": deleted, "batch_id": batch_id}
